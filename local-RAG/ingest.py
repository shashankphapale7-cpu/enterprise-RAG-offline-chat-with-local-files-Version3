"""
AI Memory OS — Ingestion Pipeline (Enterprise Edition)
Reads files from data/<tenant>/ (recursively, including nested folders),
extracts text, chunks, embeds, and stores in ChromaDB.
Supports: txt, pdf, docx, xlsx, xls, csv, md, json, xml, html, rtf, log, pptx, images

Enterprise features:
  - Recursive nested folder traversal (rglob)
  - Streaming batch processing for millions of files (never holds all data in RAM)
  - Configurable batch sizes for files, embeddings, and ChromaDB upserts
  - Explicit garbage collection between batches
  - Tenant isolation via metadata
"""

import gc
import os
import sys
import csv
import json
import hashlib
import datetime
from pathlib import Path
from html.parser import HTMLParser
from typing import Generator

import chromadb
from rich.console import Console
from rich.progress import Progress, SpinnerColumn, TextColumn, BarColumn, TimeElapsedColumn
from embedding_utils import load_embedding_model, EMBEDDING_MODEL

# ─── Config ───────────────────────────────────────────────────────────
DATA_DIR = Path(__file__).parent / "data"
CHROMA_DIR = Path(__file__).parent / "chroma_db"
COLLECTION_NAME = "memories"
CHUNK_SIZE = 500
EXCEL_ROWS_PER_CHUNK = 15  # Rows per chunk for structured spreadsheet data

# ─── Memory Management Config ────────────────────────────────────────
FILE_BATCH_SIZE = 50       # Process this many files before flushing to ChromaDB
EMBED_BATCH_SIZE = 64      # Sentences per embedding batch (SentenceTransformer internal)
UPSERT_BATCH_SIZE = 500    # Max chunks per ChromaDB upsert call

# Supported file extensions
TEXT_EXTENSIONS = {".txt", ".md", ".json", ".xml", ".rtf", ".log"}
DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".pptx"}
SPREADSHEET_EXTENSIONS = {".xlsx", ".xls", ".csv"}
WEB_EXTENSIONS = {".html", ".htm"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tiff", ".webp"}

ALL_SUPPORTED = TEXT_EXTENSIONS | DOCUMENT_EXTENSIONS | SPREADSHEET_EXTENSIONS | WEB_EXTENSIONS | IMAGE_EXTENSIONS

console = Console()


# ─── HTML Text Stripper ──────────────────────────────────────────────
class HTMLTextExtractor(HTMLParser):
    """Strip HTML tags and extract plain text."""
    def __init__(self):
        super().__init__()
        self._result = []
        self._skip = False

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self._skip = True

    def handle_endtag(self, tag):
        if tag in ("script", "style"):
            self._skip = False

    def handle_data(self, data):
        if not self._skip:
            self._result.append(data)

    def get_text(self):
        return " ".join(self._result)


# ─── System / Hidden Dirs to Ignore ───────────────────────────────────
IGNORED_DIRS = {
    ".git", ".venv", "venv", "chroma_db", "node_modules",
    "__pycache__", ".agents", ".streamlit", ".gemini", ".idea", ".vscode"
}


# ─── File Discovery (Generator — never builds full list in memory) ───
def discover_files(root_dir: Path) -> Generator[Path, None, None]:
    """
    Recursively yield all supported files under root_dir.
    Uses a generator so we never hold millions of Path objects in RAM.
    Skips hidden/system directories to avoid recursive loops.
    """
    for filepath in root_dir.rglob("*"):
        if any(part.startswith(".") or part in IGNORED_DIRS for part in filepath.parts):
            continue
        if filepath.is_file() and filepath.suffix.lower() in ALL_SUPPORTED:
            yield filepath


def count_files_fast(root_dir: Path) -> int:
    """
    Count supported files without building a list.
    For very large directories, this is the quick pre-scan.
    """
    count = 0
    for filepath in root_dir.rglob("*"):
        if any(part.startswith(".") or part in IGNORED_DIRS for part in filepath.parts):
            continue
        if filepath.is_file() and filepath.suffix.lower() in ALL_SUPPORTED:
            count += 1
    return count



# ─── File Readers ─────────────────────────────────────────────────────
def read_text_file(filepath: Path) -> str:
    """Read plain text files."""
    try:
        encodings = ["utf-8", "latin-1", "cp1252"]
        for enc in encodings:
            try:
                return filepath.read_text(encoding=enc)
            except UnicodeDecodeError:
                continue
        # Fallback: read as bytes and decode with errors='replace'
        return filepath.read_bytes().decode("utf-8", errors="replace")
    except Exception as e:
        console.print(f"  [red]Error reading {filepath.name}: {e}[/red]")
        return ""


def read_pdf(filepath: Path) -> str:
    """Extract text from PDF using pypdf."""
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(filepath))
        text_parts = []
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
        return "\n".join(text_parts)
    except Exception as e:
        console.print(f"  [red]Error reading PDF {filepath.name}: {e}[/red]")
        return ""


def read_docx(filepath: Path) -> str:
    """Extract text from .docx using python-docx."""
    try:
        from docx import Document
        doc = Document(str(filepath))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also read tables
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells)
                if row_text.strip():
                    paragraphs.append(row_text)
        return "\n".join(paragraphs)
    except Exception as e:
        console.print(f"  [red]Error reading DOCX {filepath.name}: {e}[/red]")
        return ""


def read_pptx(filepath: Path) -> str:
    """Extract text from .pptx using python-pptx."""
    try:
        from pptx import Presentation
        prs = Presentation(str(filepath))
        text_parts = []
        for slide_num, slide in enumerate(prs.slides, 1):
            slide_texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_texts.append(shape.text.strip())
            if slide_texts:
                text_parts.append(f"[Slide {slide_num}]\n" + "\n".join(slide_texts))
        return "\n\n".join(text_parts)
    except Exception as e:
        console.print(f"  [red]Error reading PPTX {filepath.name}: {e}[/red]")
        return ""


def _format_cell_value(val) -> str:
    """Format a cell value for text representation, handling dates, numbers, etc."""
    if val is None:
        return ""
    if isinstance(val, datetime.datetime):
        return val.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(val, datetime.date):
        return val.strftime("%Y-%m-%d")
    if isinstance(val, float):
        # Avoid scientific notation for large numbers
        if val == int(val):
            return str(int(val))
        return f"{val:.4f}".rstrip("0").rstrip(".")
    return str(val)


def read_xlsx(filepath: Path) -> list[str]:
    """
    Extract structured chunks from .xlsx files.
    Returns a list of pre-formed chunks (not a single text string).
    Each chunk contains the sheet name, header row, and a batch of data rows
    so that context is preserved within every chunk.
    """
    chunks = []
    try:
        from openpyxl import load_workbook
        wb = load_workbook(str(filepath), read_only=True, data_only=True)

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue

            # Extract headers from first row
            headers = [str(h).strip() if h is not None else f"Column_{i+1}"
                       for i, h in enumerate(rows[0])]
            header_line = " | ".join(headers)

            # Process data rows in batches
            data_rows = rows[1:]
            if not data_rows:
                # Sheet with only headers
                chunks.append(
                    f"[File: {filepath.name} | Sheet: {sheet_name}]\n"
                    f"Headers: {header_line}\n"
                    f"(Empty sheet — no data rows)"
                )
                continue

            for batch_start in range(0, len(data_rows), EXCEL_ROWS_PER_CHUNK):
                batch_end = min(batch_start + EXCEL_ROWS_PER_CHUNK, len(data_rows))
                batch = data_rows[batch_start:batch_end]

                chunk_lines = [
                    f"[File: {filepath.name} | Sheet: {sheet_name}]",
                    f"Headers: {header_line}",
                    f"Rows {batch_start + 1}-{batch_end} of {len(data_rows)}:",
                    ""
                ]

                for row_idx, row in enumerate(batch, start=batch_start + 1):
                    row_parts = []
                    for i, val in enumerate(row):
                        formatted = _format_cell_value(val)
                        if formatted:
                            header = headers[i] if i < len(headers) else f"Column_{i+1}"
                            row_parts.append(f"{header}: {formatted}")
                    if row_parts:
                        chunk_lines.append(f"  Row {row_idx}: {', '.join(row_parts)}")

                if len(chunk_lines) > 4:  # Has actual data beyond header
                    chunks.append("\n".join(chunk_lines))

        wb.close()
    except Exception as e:
        console.print(f"  [red]Error reading XLSX {filepath.name}: {e}[/red]")

    return chunks


def read_xls(filepath: Path) -> list[str]:
    """
    Extract structured chunks from legacy .xls files using xlrd.
    Returns pre-formed chunks like read_xlsx.
    """
    chunks = []
    try:
        import xlrd
        wb = xlrd.open_workbook(str(filepath))

        for sheet_idx in range(wb.nsheets):
            ws = wb.sheet_by_index(sheet_idx)
            sheet_name = ws.name
            if ws.nrows == 0:
                continue

            # Extract headers from first row
            headers = [str(ws.cell_value(0, c)).strip() or f"Column_{c+1}"
                       for c in range(ws.ncols)]
            header_line = " | ".join(headers)

            data_row_count = ws.nrows - 1
            if data_row_count <= 0:
                chunks.append(
                    f"[File: {filepath.name} | Sheet: {sheet_name}]\n"
                    f"Headers: {header_line}\n"
                    f"(Empty sheet — no data rows)"
                )
                continue

            for batch_start in range(0, data_row_count, EXCEL_ROWS_PER_CHUNK):
                batch_end = min(batch_start + EXCEL_ROWS_PER_CHUNK, data_row_count)

                chunk_lines = [
                    f"[File: {filepath.name} | Sheet: {sheet_name}]",
                    f"Headers: {header_line}",
                    f"Rows {batch_start + 1}-{batch_end} of {data_row_count}:",
                    ""
                ]

                for row_idx in range(batch_start + 1, batch_end + 1):
                    row_parts = []
                    for c in range(ws.ncols):
                        val = ws.cell_value(row_idx, c)
                        formatted = _format_cell_value(val)
                        if formatted:
                            header = headers[c] if c < len(headers) else f"Column_{c+1}"
                            row_parts.append(f"{header}: {formatted}")
                    if row_parts:
                        chunk_lines.append(f"  Row {row_idx}: {', '.join(row_parts)}")

                if len(chunk_lines) > 4:
                    chunks.append("\n".join(chunk_lines))

    except ImportError:
        console.print(f"  [yellow]xlrd not installed — cannot read .xls files. Install with: pip install xlrd[/yellow]")
    except Exception as e:
        console.print(f"  [red]Error reading XLS {filepath.name}: {e}[/red]")

    return chunks


def read_csv_file(filepath: Path) -> list[str]:
    """
    Extract structured chunks from CSV files.
    Returns pre-formed chunks with header context preserved per chunk.
    """
    chunks = []
    try:
        with open(filepath, "r", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f)
            rows = list(reader)
            if not rows:
                return []

            headers = [h.strip() or f"Column_{i+1}" for i, h in enumerate(rows[0])]
            header_line = " | ".join(headers)
            data_rows = rows[1:]

            if not data_rows:
                chunks.append(
                    f"[File: {filepath.name}]\n"
                    f"Headers: {header_line}\n"
                    f"(Empty file — no data rows)"
                )
                return chunks

            for batch_start in range(0, len(data_rows), EXCEL_ROWS_PER_CHUNK):
                batch_end = min(batch_start + EXCEL_ROWS_PER_CHUNK, len(data_rows))
                batch = data_rows[batch_start:batch_end]

                chunk_lines = [
                    f"[File: {filepath.name}]",
                    f"Headers: {header_line}",
                    f"Rows {batch_start + 1}-{batch_end} of {len(data_rows)}:",
                    ""
                ]

                for row_idx, row in enumerate(batch, start=batch_start + 1):
                    row_parts = []
                    for i, val in enumerate(row):
                        if val.strip():
                            header = headers[i] if i < len(headers) else f"Column_{i+1}"
                            row_parts.append(f"{header}: {val.strip()}")
                    if row_parts:
                        chunk_lines.append(f"  Row {row_idx}: {', '.join(row_parts)}")

                if len(chunk_lines) > 4:
                    chunks.append("\n".join(chunk_lines))

    except Exception as e:
        console.print(f"  [red]Error reading CSV {filepath.name}: {e}[/red]")

    return chunks


def read_html(filepath: Path) -> str:
    """Extract text from HTML files."""
    try:
        raw = filepath.read_text(encoding="utf-8", errors="replace")
        extractor = HTMLTextExtractor()
        extractor.feed(raw)
        return extractor.get_text()
    except Exception as e:
        console.print(f"  [red]Error reading HTML {filepath.name}: {e}[/red]")
        return ""


def read_image(filepath: Path) -> str:
    """Extract text from images using OCR (pytesseract) and return description."""
    text_parts = []
    text_parts.append(f"[Image: {filepath.name}]")

    # Try OCR
    try:
        from PIL import Image
        import pytesseract
        img = Image.open(filepath)
        ocr_text = pytesseract.image_to_string(img).strip()
        if ocr_text:
            text_parts.append(f"OCR Text: {ocr_text}")
        img.close()
    except ImportError:
        text_parts.append("(OCR not available — install pytesseract and Tesseract engine)")
    except Exception as e:
        text_parts.append(f"(OCR failed: {e})")

    # Basic image metadata
    try:
        from PIL import Image
        img = Image.open(filepath)
        w, h = img.size
        text_parts.append(f"Dimensions: {w}x{h}, Format: {img.format}, Mode: {img.mode}")
        img.close()
    except Exception:
        pass

    return "\n".join(text_parts)


def extract_text(filepath: Path) -> str | list[str]:
    """
    Route file to the appropriate reader based on extension.
    Returns either a string (for text-based files) or a list of pre-formed chunks
    (for spreadsheet files that need structured chunking).
    """
    ext = filepath.suffix.lower()

    if ext in TEXT_EXTENSIONS:
        return read_text_file(filepath)
    elif ext == ".pdf":
        return read_pdf(filepath)
    elif ext == ".docx":
        return read_docx(filepath)
    elif ext == ".pptx":
        return read_pptx(filepath)
    elif ext == ".xlsx":
        return read_xlsx(filepath)
    elif ext == ".xls":
        return read_xls(filepath)
    elif ext == ".csv":
        return read_csv_file(filepath)
    elif ext in WEB_EXTENSIONS:
        return read_html(filepath)
    elif ext in IMAGE_EXTENSIONS:
        return read_image(filepath)
    else:
        return ""


# ─── Chunking ─────────────────────────────────────────────────────────
def chunk_text(text: str, chunk_size: int = CHUNK_SIZE) -> list[str]:
    """Split text into chunks of approximately `chunk_size` characters."""
    if not text.strip():
        return []
    chunks = []
    words = text.split()
    current_chunk = []
    current_len = 0
    for word in words:
        word_len = len(word) + 1  # +1 for space
        if current_len + word_len > chunk_size and current_chunk:
            chunks.append(" ".join(current_chunk))
            current_chunk = [word]
            current_len = word_len
        else:
            current_chunk.append(word)
            current_len += word_len
    if current_chunk:
        chunks.append(" ".join(current_chunk))
    return chunks


# ─── Unique ID Generation ────────────────────────────────────────────
def make_chunk_id(relative_path: str, chunk_index: int, tenant_id: str = "") -> str:
    """
    Generate a deterministic, unique ID for each chunk, scoped to tenant.
    Uses relative path (not just filename) to handle nested folder structures.
    """
    raw = f"{tenant_id}::{relative_path}::chunk_{chunk_index}"
    return hashlib.md5(raw.encode()).hexdigest()


# ─── Batch Flush (memory-safe upsert) ────────────────────────────────
def _flush_batch(
    collection,
    batch_ids: list[str],
    batch_docs: list[str],
    batch_embeddings: list[list[float]],
    batch_metadatas: list[dict],
):
    """
    Upsert a batch of chunks into ChromaDB, respecting the max upsert size.
    After upserting, the caller should clear the lists and call gc.collect().
    """
    if not batch_ids:
        return

    for i in range(0, len(batch_ids), UPSERT_BATCH_SIZE):
        end = min(i + UPSERT_BATCH_SIZE, len(batch_ids))
        collection.upsert(
            ids=batch_ids[i:end],
            documents=batch_docs[i:end],
            embeddings=batch_embeddings[i:end],
            metadatas=batch_metadatas[i:end],
        )


# ─── Main Ingestion (Streaming, Memory-Safe) ─────────────────────────
def ingest(
    data_dir: Path = DATA_DIR,
    show_progress: bool = True,
    tenant_id: str = "",
    progress_callback=None,
) -> dict:
    """
    Main ingestion pipeline — enterprise grade.

    Architecture for millions of files:
      1. Generator-based file discovery (rglob) — never holds full file list in RAM
      2. Files processed in batches of FILE_BATCH_SIZE
      3. After each batch: embeddings generated, upserted to ChromaDB, memory freed
      4. Explicit gc.collect() between batches to reclaim memory

    Args:
        data_dir: Root directory to scan (recursively).
        show_progress: Whether to show rich progress bar (CLI mode).
        tenant_id: User/tenant identifier for data isolation.
        progress_callback: Optional callable(files_done, total_files, current_file)
                           for Streamlit progress updates.
    Returns:
        dict with stats: {files_processed, chunks_created, files_skipped, errors}
    """
    if not data_dir.exists():
        data_dir.mkdir(parents=True)
        console.print(f"[yellow]Created data directory: {data_dir}[/yellow]")
        return {"files_processed": 0, "chunks_created": 0, "files_skipped": 0, "errors": []}

    # ── Phase 1: Quick file count scan ────────────────────────────────
    console.print(f"\n[bold cyan]🧠 AI Memory OS — Ingestion Pipeline[/bold cyan]")
    console.print(f"[dim]Scanning {data_dir} recursively...[/dim]")

    total_files = count_files_fast(data_dir)
    if total_files == 0:
        console.print("[yellow]No supported files found.[/yellow]")
        return {"files_processed": 0, "chunks_created": 0, "files_skipped": 0, "errors": []}

    console.print(f"[dim]Found {total_files:,} supported file(s) for tenant: {tenant_id or 'default'}[/dim]\n")

    # ── Phase 2: Load embedding model (cached globally) ──────────────
    console.print("[cyan]Loading embedding model...[/cyan]")
    model = load_embedding_model(EMBEDDING_MODEL)
    console.print(f"[green]✓ Model loaded: {EMBEDDING_MODEL}[/green]\n")

    # ── Phase 3: Initialize ChromaDB ─────────────────────────────────
    client = chromadb.PersistentClient(path=str(CHROMA_DIR))
    collection = client.get_or_create_collection(
        name=COLLECTION_NAME,
        metadata={"hnsw:space": "cosine"}
    )

    # ── Phase 4: Streaming batch processing ──────────────────────────
    stats = {"files_processed": 0, "chunks_created": 0, "files_skipped": 0, "errors": []}

    # Accumulators for current batch (cleared after each flush)
    batch_ids = []
    batch_docs = []
    batch_embeddings = []
    batch_metadatas = []
    files_in_current_batch = 0

    with Progress(
        SpinnerColumn(),
        TextColumn("[progress.description]{task.description}"),
        BarColumn(),
        TextColumn("{task.completed:,}/{task.total:,}"),
        TimeElapsedColumn(),
        console=console,
        disable=not show_progress,
    ) as progress:
        task = progress.add_task("Processing files...", total=total_files)

        for filepath in discover_files(data_dir):
            # Compute relative path for nested folder support
            try:
                rel_path = str(filepath.relative_to(data_dir))
            except ValueError:
                rel_path = filepath.name

            progress.update(task, description=f"Processing {rel_path}...")

            # Notify Streamlit callback if provided
            if progress_callback:
                progress_callback(stats["files_processed"], total_files, rel_path)

            # ── Extract text ──────────────────────────────────────────
            try:
                extracted = extract_text(filepath)
            except Exception as e:
                stats["errors"].append(f"{rel_path}: Extract failed — {e}")
                stats["files_skipped"] += 1
                progress.advance(task)
                continue

            # ── Convert to chunks ─────────────────────────────────────
            if isinstance(extracted, list):
                # Spreadsheet: already pre-chunked
                chunks = extracted
                if not chunks:
                    stats["errors"].append(f"{rel_path}: No data extracted")
                    stats["files_skipped"] += 1
                    progress.advance(task)
                    continue
            else:
                # Plain text: chunk it
                text = extracted
                if not text or not text.strip():
                    stats["errors"].append(f"{rel_path}: No text extracted")
                    stats["files_skipped"] += 1
                    progress.advance(task)
                    continue
                chunks = chunk_text(text)
                if not chunks:
                    stats["errors"].append(f"{rel_path}: No chunks created")
                    stats["files_skipped"] += 1
                    progress.advance(task)
                    continue

            # ── Generate embeddings (batched internally by SentenceTransformer) ──
            embeddings = model.encode(
                chunks,
                batch_size=EMBED_BATCH_SIZE,
                show_progress_bar=False,
            ).tolist()

            # ── Build metadata and accumulate ─────────────────────────
            timestamp = datetime.datetime.now().isoformat()
            for i, (chunk, embedding) in enumerate(zip(chunks, embeddings)):
                chunk_id = make_chunk_id(rel_path, i, tenant_id=tenant_id)
                metadata = {
                    "source": rel_path,
                    "file_type": filepath.suffix.lower(),
                    "chunk_index": i,
                    "total_chunks": len(chunks),
                    "indexed_at": timestamp,
                }
                if tenant_id:
                    metadata["tenant_id"] = tenant_id

                batch_ids.append(chunk_id)
                batch_docs.append(chunk)
                batch_embeddings.append(embedding)
                batch_metadatas.append(metadata)

            stats["files_processed"] += 1
            stats["chunks_created"] += len(chunks)
            files_in_current_batch += 1

            # Free extracted text / chunks from this file immediately
            del extracted, chunks, embeddings

            # ── Flush batch when threshold reached ────────────────────
            if files_in_current_batch >= FILE_BATCH_SIZE:
                _flush_batch(collection, batch_ids, batch_docs, batch_embeddings, batch_metadatas)

                # Clear accumulators and reclaim memory
                batch_ids.clear()
                batch_docs.clear()
                batch_embeddings.clear()
                batch_metadatas.clear()
                files_in_current_batch = 0
                gc.collect()

            progress.advance(task)

        # ── Flush remaining chunks ────────────────────────────────────
        _flush_batch(collection, batch_ids, batch_docs, batch_embeddings, batch_metadatas)
        batch_ids.clear()
        batch_docs.clear()
        batch_embeddings.clear()
        batch_metadatas.clear()
        gc.collect()

    # ── Summary ───────────────────────────────────────────────────────
    console.print(f"\n[bold green]✅ Ingestion Complete![/bold green]")
    console.print(f"  📁 Files processed: [cyan]{stats['files_processed']:,}[/cyan]")
    console.print(f"  🧩 Chunks created:  [cyan]{stats['chunks_created']:,}[/cyan]")
    console.print(f"  ⏭️  Files skipped:   [cyan]{stats['files_skipped']:,}[/cyan]")
    console.print(f"  💾 Collection size: [cyan]{collection.count():,}[/cyan] total memories")

    if stats["errors"]:
        console.print(f"\n  [yellow]⚠ Warnings ({len(stats['errors'])}):[/yellow]")
        # Show at most 20 errors to avoid flooding the console
        for err in stats["errors"][:20]:
            console.print(f"    [dim]{err}[/dim]")
        if len(stats["errors"]) > 20:
            console.print(f"    [dim]... and {len(stats['errors']) - 20} more[/dim]")

    return stats


if __name__ == "__main__":
    ingest()
